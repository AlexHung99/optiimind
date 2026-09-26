"""Bounded local document skills for OptiChat Agent."""
from pathlib import Path
from uuid import uuid4

from opti_pdf import read_pdf


MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
TEXT_TYPES = {'.txt', '.md', '.markdown', '.csv'}


def _check(path, suffixes):
    if path.suffix.lower() not in suffixes or not path.is_file():
        raise ValueError('檔案類型或路徑不支援。')
    if path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError('文件超過 20 MB，請先縮小檔案。')


def _output(path):
    # A fresh sibling keeps the original untouched and stays in the selected folder.
    return path.with_name(f'{path.stem}-OptiChat-{uuid4().hex[:8]}{path.suffix}')


def read_document(path, kind, page=None):
    if kind == 'read_pdf':
        document = read_pdf(path)
        text = document['text']
        if page is not None:
            number = int(page)
            if not 1 <= number <= document['pages']:
                raise ValueError('PDF 頁碼超出範圍。')
            marker = f'\n[第 {number} 頁]\n'
            text = text.split(marker, 1)[1].split('\n[第 ', 1)[0] if marker in text else ''
        return (f"PDF 共 {document['pages']} 頁。" +
                (' 僅擷取到部分內容。' if document['truncated'] else '') + '\n' +
                (text[:3000] if text else '沒有可擷取的文字層；掃描 PDF 請在聊天中選頁使用圖片模型。'))
    if kind == 'read_excel':
        _check(path, {'.xlsx'})
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheet_name = page or book.sheetnames[0]
            if sheet_name not in book.sheetnames:
                raise ValueError('找不到指定的工作表。')
            sheet = book[sheet_name]
            lines = [f"工作表：{', '.join(book.sheetnames[:20])}",
                     f'目前預覽：{sheet_name}，{sheet.max_row} 列 × {sheet.max_column} 欄']
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25),
                                       max_col=min(sheet.max_column, 12)):
                values = [f'{cell.coordinate}={str(cell.value)[:80]}' for cell in row if cell.value is not None]
                if values:
                    lines.append('  '.join(values))
            return '\n'.join(lines)[:3200]
        finally:
            book.close()
    if kind == 'read_word':
        _check(path, {'.docx'})
        from docx import Document
        document = Document(path)
        lines = [f'段落 {i+1}：{p.text}' for i, p in enumerate(document.paragraphs[:80]) if p.text.strip()]
        for table_number, table in enumerate(document.tables[:5], 1):
            for row in table.rows[:15]:
                lines.append(f'表格 {table_number}：' + ' | '.join(cell.text[:100] for cell in row.cells[:10]))
        return '\n'.join(lines)[:3200] or 'Word 文件沒有可預覽的文字。'
    raise ValueError('不支援的文件讀取工具。')


def prepare_edit(path, action):
    """Validate and capture an edit. Returns a preview and a one-shot save function."""
    kind = action['action']
    target = _output(path)
    if kind == 'edit_text':
        _check(path, TEXT_TYPES)
    elif kind == 'edit_excel':
        _check(path, {'.xlsx'})
    elif kind == 'edit_word':
        _check(path, {'.docx'})
    else:
        raise ValueError('不支援的編輯工具。')
    original = path.read_bytes()
    if kind == 'edit_text':
        find = str(action.get('find') or '')
        replace = str(action.get('replace') or '')
        if not find or len(find) > 1000 or len(replace) > 3000:
            raise ValueError('請提供長度適當的原文和替換文字。')
        source = original.decode('utf-8-sig')
        if source.count(find) != 1:
            raise ValueError('原文必須恰好出現一次；請提供更明確的文字。')
        result = source.replace(find, replace, 1).encode('utf-8')
        preview = f'文字替換\n原文：{find[:180]}\n改為：{replace[:180]}'
        def write():
            with target.open('xb') as output:
                output.write(result)
    elif kind == 'edit_excel':
        from openpyxl import load_workbook
        sheet_name = str(action.get('sheet') or '').strip()
        edits = action.get('cells')
        if not sheet_name or not isinstance(edits, list) or not 1 <= len(edits) <= 30:
            raise ValueError('請指定工作表與 1 至 30 個儲存格。')
        book = load_workbook(path, keep_links=False)
        if sheet_name not in book.sheetnames:
            raise ValueError('找不到指定的工作表。')
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        sheet = book[sheet_name]
        lines = []
        for edit in edits:
            if not isinstance(edit, dict) or not isinstance(edit.get('cell'), str):
                raise ValueError('儲存格資料格式無效。')
            cell = edit['cell'].upper()
            try:
                column, row = coordinate_from_string(cell)
                if row > 100000 or column_index_from_string(column) > 1000:
                    raise ValueError()
            except ValueError as error:
                raise ValueError('儲存格座標無效或超出可編輯範圍。') from error
            value = edit.get('value')
            if not isinstance(value, (str, int, float, bool, type(None))) or isinstance(value, str) and len(value) > 2000:
                raise ValueError('儲存格內容只支援短文字、數字、布林值或空值。')
            previous = sheet[cell].value
            sheet[cell] = value
            lines.append(f'{cell}: {str(previous)[:60]} → {str(value)[:60]}')
        preview = 'Excel · ' + sheet_name + '\n' + '\n'.join(lines)
        def write():
            with target.open('xb') as output:
                book.save(output)
    elif kind == 'edit_word':
        from docx import Document
        document = Document(path)
        operation = action.get('operation')
        if operation == 'append':
            value = str(action.get('text') or '')
            if not value or len(value) > 3000:
                raise ValueError('請提供 1 至 3000 字的新段落。')
            document.add_paragraph(value)
            preview = 'Word · 新增段落\n' + value[:400]
        elif operation == 'replace':
            find = str(action.get('find') or '')
            replace = str(action.get('replace') or '')
            if not find or len(find) > 1000 or len(replace) > 3000:
                raise ValueError('請提供長度適當的原文和替換文字。')
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
            runs = [run for paragraph in paragraphs for run in paragraph.runs if find in run.text]
            if len(runs) != 1 or runs[0].text.count(find) != 1:
                raise ValueError('原文須恰好位於一個格式片段中；跨格式文字請改用新增段落。')
            runs[0].text = runs[0].text.replace(find, replace, 1)
            preview = f'Word · 替換文字\n原文：{find[:180]}\n改為：{replace[:180]}'
        else:
            raise ValueError('Word 操作只支援 append 或 replace。')
        def write():
            with target.open('xb') as output:
                document.save(output)
    def apply():
        if path.read_bytes() != original:
            raise ValueError('確認期間原檔已改變，請重新讀取後再試。')
        write()
        return '已另存新檔：' + str(target)
    return preview + '\n原檔保留，將另存：' + str(target), apply

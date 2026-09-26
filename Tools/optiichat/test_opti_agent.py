"""Agent loop and UI regression checks without calling an actual model."""
from contextlib import ExitStack
from copy import deepcopy
from pathlib import Path
import tempfile
import time
import tkinter as tk
import unittest
from unittest.mock import patch

import opti_agent
import opti_agent_ui
import opti_app
import opti_core
from opti_document import read_document


class AgentCoreTests(unittest.TestCase):
    def test_document_skills_keep_original_and_require_approval(self):
        from docx import Document
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            note = root/'note.txt'
            note.write_text('舊內容', encoding='utf-8')
            workbook = Workbook()
            workbook.active.title = '資料'
            workbook.active['A1'] = '舊值'
            workbook.save(root/'data.xlsx')
            word = Document()
            word.add_paragraph('原本段落')
            word.save(root/'report.docx')
            tools = opti_agent.LocalTools(root)
            self.assertIn('A1=舊值', tools.run({'action': 'read_excel', 'path': 'data.xlsx'}))
            self.assertIn('原本段落', tools.run({'action': 'read_word', 'path': 'report.docx'}))
            cases = [
                ({'action': 'edit_text', 'path': 'note.txt', 'find': '舊內容', 'replace': '新內容'}, note),
                ({'action': 'edit_excel', 'path': 'data.xlsx', 'sheet': '資料',
                  'cells': [{'cell': 'A1', 'value': '新值'}]}, root/'data.xlsx'),
                ({'action': 'edit_word', 'path': 'report.docx', 'operation': 'replace',
                  'find': '原本', 'replace': '更新'}, root/'report.docx'),
            ]
            for action, source in cases:
                before = source.read_bytes()
                preview, apply = tools.plan(action)
                self.assertIn('另存', preview)
                self.assertEqual(source.read_bytes(), before)
                result = apply()
                output = Path(result.split('：', 1)[1])
                self.assertTrue(output.is_file())
                self.assertEqual(source.read_bytes(), before)
                if output.suffix == '.txt':
                    self.assertEqual(output.read_text(encoding='utf-8'), '新內容')
                elif output.suffix == '.xlsx':
                    self.assertEqual(load_workbook(output).active['A1'].value, '新值')
                else:
                    self.assertIn('更新段落', Document(output).paragraphs[0].text)
            with self.assertRaises(ValueError):
                tools.plan({'action': 'edit_text', 'path': '../note.txt', 'find': 'a', 'replace': 'b'})

    def test_agent_denied_edit_does_not_write(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'note.txt'
            path.write_text('原文', encoding='utf-8')
            replies = iter(['{"action":"edit_text","path":"note.txt","find":"原文","replace":"新文"}',
                            '{"action":"finish","answer":"編輯未獲確認。"}'])
            runner = opti_agent.AgentRunner(temporary, lambda messages, cancel: next(replies),
                                            lambda event, value: None,
                                            lambda preview, cancel: False)
            self.assertIn('未獲確認', runner.run('修改文字'))
            self.assertEqual(list(Path(temporary).iterdir()), [path])
            self.assertEqual(path.read_text(encoding='utf-8'), '原文')

    def test_agent_pdf_reads_text_without_edit_action(self):
        from test_opti_chat import PdfTests
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'report.pdf'
            PdfTests.text_pdf(path)
            tools = opti_agent.LocalTools(temporary)
            self.assertIn('Hello PDF document', tools.run({'action': 'read_pdf', 'path': 'report.pdf', 'page': 1}))
            with self.assertRaises(ValueError):
                opti_agent.parse_action('{"action":"edit_pdf","path":"report.pdf"}')

    def test_folder_tools_reject_traversal_and_keep_reads_bounded(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root/'notes.txt').write_text('甲：測試內容\n乙：第二行', encoding='utf-8')
            (root/'.env').write_text('SECRET', encoding='utf-8')
            (root/'large.txt').write_text('x' * (opti_agent.MAX_FILE_BYTES + 1), encoding='utf-8')
            tool = opti_agent.LocalTools(root)
            self.assertIn('notes.txt', tool.run({'action': 'list_files', 'path': '.'}))
            self.assertNotIn('.env', tool.run({'action': 'list_files', 'path': '.'}))
            self.assertIn('測試內容', tool.run({'action': 'read_file', 'path': 'notes.txt'}))
            self.assertIn('notes.txt:1', tool.run({'action': 'search_files', 'query': '測試', 'path': '.'}))
            with self.assertRaises(ValueError):
                tool.run({'action': 'read_file', 'path': '../outside.txt'})
            with self.assertRaises(ValueError):
                tool.run({'action': 'read_file', 'path': '.env'})
            with self.assertRaises(ValueError):
                tool.run({'action': 'read_file', 'path': 'large.txt'})

    def test_agent_uses_tool_result_then_finishes(self):
        with tempfile.TemporaryDirectory() as temporary:
            (Path(temporary)/'notes.txt').write_text('資料：42', encoding='utf-8')
            events = []
            replies = iter(['{"action":"read_file","path":"notes.txt"}',
                            '```json\n{"action":"finish","answer":"答案是 42。"}\n```'])
            seen = []
            def model_call(messages, cancelled):
                seen.append(messages)
                return next(replies)
            runner = opti_agent.AgentRunner(temporary, model_call,
                                            lambda event, value: events.append((event, value)))
            self.assertEqual(runner.run('讀取 notes.txt'), '答案是 42。')
            self.assertEqual([event for event, _ in events if event == 'step'], ['step'])
            self.assertIn('資料：42', seen[-1][-1]['content'])

    def test_task_store_is_separate_and_validates_id(self):
        with tempfile.TemporaryDirectory() as temporary:
            task = opti_agent.new_task()
            task['prompt'] = '整理檔案'
            opti_agent.save_task(task, temporary)
            self.assertEqual(opti_agent.load_task(task['id'], temporary)['prompt'], '整理檔案')
            self.assertEqual(len(opti_agent.list_tasks(temporary)), 1)
            with self.assertRaises(ValueError):
                opti_agent.task_path('../bad', temporary)
            opti_agent.delete_task(task['id'], temporary)
            self.assertEqual(opti_agent.list_tasks(temporary), [])

    def test_cancel_prevents_a_late_model_result_from_becoming_a_final_answer(self):
        events = []
        runner = None
        def model_call(messages, cancelled):
            runner.cancel()
            return '{"action":"finish","answer":"不應出現"}'
        runner = opti_agent.AgentRunner('', model_call,
                                        lambda event, value: events.append((event, value)))
        self.assertIsNone(runner.run('取消測試'))
        self.assertFalse(any(event == 'final' for event, _ in events))


class AgentUiTests(unittest.TestCase):
    def test_document_edit_confirmed_through_agent_ui(self):
        with ExitStack() as stack:
            store = Path(stack.enter_context(tempfile.TemporaryDirectory()))
            folder = Path(stack.enter_context(tempfile.TemporaryDirectory()))
            source = folder/'draft.txt'
            source.write_text('舊句', encoding='utf-8')
            stack.enter_context(patch.object(opti_app, 'load_settings', return_value=deepcopy(opti_core.DEFAULTS)))
            for method in ('start_tray', 'monitor', 'connect', 'check_updates'):
                stack.enter_context(patch.object(opti_app.OptiiApp, method))
            stack.enter_context(patch.object(opti_agent_ui, 'list_tasks', side_effect=lambda: opti_agent.list_tasks(store)))
            stack.enter_context(patch.object(opti_agent_ui, 'save_task', side_effect=lambda task: opti_agent.save_task(task, store)))
            confirmed = stack.enter_context(patch.object(opti_app.OptiiApp, 'confirm', return_value=True))
            root = tk.Tk()
            app = opti_app.OptiiApp(root)
            try:
                workspace = app.agent_workspace
                workspace.folder.set(str(folder))
                workspace.goal.insert('1.0', '修改 draft.txt')
                app.catalog = [{'name': opti_core.DEFAULT_MODEL,
                                'capabilities': ['completion', 'vision'],
                                'details': {'family': 'mllama'}}]
                app.ready = True
                app.model_loading = False
                replies = iter(['{"action":"edit_text","path":"draft.txt","find":"舊句","replace":"新句"}',
                                '{"action":"finish","answer":"已另存修訂版。"}'])
                workspace._model_call = lambda messages, cancelled: next(replies)
                workspace.start()
                limit = time.monotonic() + 5
                while workspace.runner and time.monotonic() < limit:
                    root.update()
                    time.sleep(.02)
                self.assertIsNone(workspace.runner)
                self.assertEqual(workspace.task['status'], 'done')
                confirmed.assert_called_once()
                self.assertEqual(source.read_text(encoding='utf-8'), '舊句')
                copies = list(folder.glob('draft-OptiChat-*.txt'))
                self.assertEqual(len(copies), 1)
                self.assertEqual(copies[0].read_text(encoding='utf-8'), '新句')
            finally:
                app.quit()

    def test_tabs_switch_and_agent_timeline_runs(self):
        with ExitStack() as stack:
            store = Path(stack.enter_context(tempfile.TemporaryDirectory()))
            stack.enter_context(patch.object(opti_app, 'load_settings', return_value=deepcopy(opti_core.DEFAULTS)))
            for method in ('start_tray', 'monitor', 'connect', 'check_updates'):
                stack.enter_context(patch.object(opti_app.OptiiApp, method))
            stack.enter_context(patch.object(opti_agent_ui, 'list_tasks', side_effect=lambda: opti_agent.list_tasks(store)))
            stack.enter_context(patch.object(opti_agent_ui, 'save_task', side_effect=lambda task: opti_agent.save_task(task, store)))
            root = tk.Tk()
            app = opti_app.OptiiApp(root)
            try:
                workspace = app.agent_workspace
                root.update()
                self.assertEqual(workspace.mode, 'chat')
                workspace.switch('agent')
                root.update()
                self.assertTrue(workspace.center.winfo_ismapped())
                self.assertFalse(workspace.chat_center.winfo_ismapped())
                with tempfile.TemporaryDirectory() as folder:
                    (Path(folder)/'notes.txt').write_text('本機內容', encoding='utf-8')
                    workspace.folder.set(folder)
                    workspace.goal.insert('1.0', '整理 notes.txt')
                    app.catalog = [{'name': opti_core.DEFAULT_MODEL,
                                    'capabilities': ['completion', 'vision'],
                                    'details': {'family': 'mllama'}}]
                    app.ready = True
                    app.model_loading = False
                    replies = iter(['{"action":"read_file","path":"notes.txt"}',
                                    '{"action":"finish","answer":"已整理本機內容。"}'])
                    workspace._model_call = lambda messages, cancelled: next(replies)
                    workspace.start()
                    limit = time.monotonic() + 5
                    while workspace.runner and time.monotonic() < limit:
                        root.update()
                        time.sleep(.02)
                    self.assertIsNone(workspace.runner)
                    self.assertEqual(workspace.task['status'], 'done')
                    self.assertIn('本機內容', workspace.output.get('1.0', 'end'))
                    self.assertIn('已整理本機內容', workspace.output.get('1.0', 'end'))
                    self.assertEqual(len(opti_agent.list_tasks(store)), 1)
                workspace.switch('chat')
                root.update()
                self.assertTrue(workspace.chat_center.winfo_ismapped())
            finally:
                app.quit()


if __name__ == '__main__':
    unittest.main()
